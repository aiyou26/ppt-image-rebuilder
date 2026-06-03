#!/usr/bin/env python3
"""
Rebuild a best-effort editable PowerPoint deck from exported slide images.

Default output is a visually faithful deck: each source image is placed on a
PowerPoint slide, optionally with OCR/API-generated editable text boxes overlaid.
"""

from __future__ import annotations

import argparse
import base64
import getpass
import importlib.util
import json
import mimetypes
import os
import platform
import re
import shutil
import subprocess
import sys
import tempfile
import urllib.error
import urllib.parse
import urllib.request
from dataclasses import dataclass
from pathlib import Path
from typing import Any, Dict, Iterable, List, Optional, Sequence, Tuple

try:
    from PIL import Image
except ImportError as exc:
    raise SystemExit("Missing dependency: Pillow. Install with: pip install Pillow") from exc

try:
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.util import Inches, Pt
except ImportError as exc:
    raise SystemExit("Missing dependency: python-pptx. Install with: pip install python-pptx") from exc

IMAGE_EXTENSIONS = {".png", ".jpg", ".jpeg", ".webp", ".bmp", ".tif", ".tiff"}
EMU_PER_INCH = 914400
DEFAULT_API_MODEL = "gpt-4.1-mini"
DEFAULT_API_PROVIDER = "openai"
DEFAULT_API_ENDPOINT_MODE = "auto"
DEFAULT_OCR_API_PROVIDER = "ocr-space"
DEFAULT_OCR_API_ENDPOINT = "https://api.ocr.space/parse/image"
DEFAULT_AZURE_VISION_API_VERSION = "2024-02-01"



@dataclass
class TextItem:
    text: str
    x: float
    y: float
    w: float
    h: float
    font_size: Optional[float] = None
    confidence: Optional[float] = None


@dataclass
class ApiConfig:
    provider: str
    api_key: str
    model: str
    base_url: Optional[str] = None
    endpoint_mode: str = DEFAULT_API_ENDPOINT_MODE


@dataclass
class DedicatedOcrConfig:
    provider: str
    api_key: str
    endpoint: Optional[str] = None
    language: Optional[str] = None


def natural_key(path: Path) -> List[Any]:
    parts = re.split(r"(\d+)", path.name.lower())
    return [int(part) if part.isdigit() else part for part in parts]


def collect_images(input_path: Path) -> List[Path]:
    if input_path.is_file():
        if input_path.suffix.lower() not in IMAGE_EXTENSIONS:
            raise SystemExit(f"Input file is not a supported image: {input_path}")
        return [input_path]
    if not input_path.exists():
        raise SystemExit(f"Input path does not exist: {input_path}")
    images = [p for p in input_path.iterdir() if p.is_file() and p.suffix.lower() in IMAGE_EXTENSIONS]
    images.sort(key=natural_key)
    if not images:
        raise SystemExit(f"No supported images found in: {input_path}")
    return images


def infer_slide_size(first_image: Path, slide_size: str) -> Tuple[int, int]:
    with Image.open(first_image) as img:
        width_px, height_px = img.size
    ratio = width_px / max(height_px, 1)

    if slide_size == "16:9":
        return Inches(13.333333), Inches(7.5)
    if slide_size == "4:3":
        return Inches(10), Inches(7.5)

    if abs(ratio - (16 / 9)) < 0.04:
        return Inches(13.333333), Inches(7.5)
    if abs(ratio - (4 / 3)) < 0.04:
        return Inches(10), Inches(7.5)

    if ratio >= 1:
        width_in = 13.333333
        height_in = width_in / ratio
    else:
        height_in = 7.5
        width_in = height_in * ratio
    return Inches(width_in), Inches(height_in)


def fit_rect(src_w: int, src_h: int, dst_w: int, dst_h: int) -> Tuple[int, int, int, int]:
    src_ratio = src_w / max(src_h, 1)
    dst_ratio = dst_w / max(dst_h, 1)
    if src_ratio > dst_ratio:
        width = dst_w
        height = int(dst_w / src_ratio)
        left = 0
        top = int((dst_h - height) / 2)
    else:
        height = dst_h
        width = int(dst_h * src_ratio)
        left = int((dst_w - width) / 2)
        top = 0
    return left, top, width, height


def faded_copy(image_path: Path, opacity: float, temp_dir: Path) -> Path:
    opacity = max(0.0, min(1.0, opacity))
    if opacity >= 0.999:
        return image_path
    with Image.open(image_path) as img:
        rgb = img.convert("RGB")
        white = Image.new("RGB", rgb.size, "white")
        blended = Image.blend(white, rgb, opacity)
        out = temp_dir / f"faded_{image_path.stem}.png"
        blended.save(out)
        return out


def load_ocr_json(path: Path) -> Dict[str, List[TextItem]]:
    data = json.loads(path.read_text(encoding="utf-8"))
    slides = data.get("slides", data if isinstance(data, list) else [])
    by_image: Dict[str, List[TextItem]] = {}
    for index, slide in enumerate(slides):
        image_name = str(slide.get("image", index + 1))
        raw_items = slide.get("items") or slide.get("text_blocks") or []
        items = text_items_from_raw(raw_items)
        by_image[image_name] = items
        by_image[Path(image_name).name] = items
        by_image[str(index + 1)] = items
    return by_image


def text_items_from_raw(raw_items: Any) -> List[TextItem]:
    items: List[TextItem] = []
    if not isinstance(raw_items, list):
        return items
    for raw in raw_items:
        if not isinstance(raw, dict):
            continue
        text = str(raw.get("text", "")).strip()
        if not text:
            continue
        items.append(
            TextItem(
                text=text,
                x=float(raw.get("x", 0)),
                y=float(raw.get("y", 0)),
                w=float(raw.get("w", raw.get("width", 0.1))),
                h=float(raw.get("h", raw.get("height", 0.05))),
                font_size=float(raw["font_size"]) if "font_size" in raw and raw["font_size"] is not None else None,
                confidence=float(raw["confidence"]) if "confidence" in raw and raw["confidence"] is not None else None,
            )
        )
    return items


def parse_conf(value: Any) -> float:
    try:
        return float(value)
    except Exception:
        return -1.0


def python_package_installed(import_name: str) -> bool:
    return importlib.util.find_spec(import_name) is not None


def command_exists(name: str) -> bool:
    return shutil.which(name) is not None


def run_command(cmd: Sequence[str], *, check: bool = False) -> subprocess.CompletedProcess[str]:
    printable = " ".join(cmd)
    print(f"Running: {printable}")
    return subprocess.run(cmd, text=True, stdout=sys.stdout, stderr=sys.stderr, check=check)


def pip_install(package: str) -> bool:
    try:
        run_command([sys.executable, "-m", "pip", "install", package], check=True)
        return True
    except Exception as exc:
        print(f"Warning: failed to install {package}: {exc}", file=sys.stderr)
        return False


def local_ocr_available() -> bool:
    return python_package_installed("pytesseract") and command_exists("tesseract")


def install_tesseract_binary() -> bool:
    system = platform.system().lower()

    if command_exists("tesseract"):
        return True

    if command_exists("mamba"):
        return run_best_effort([["mamba", "install", "-y", "-c", "conda-forge", "tesseract"]])
    if command_exists("conda"):
        return run_best_effort([["conda", "install", "-y", "-c", "conda-forge", "tesseract"]])

    if system == "darwin" and command_exists("brew"):
        return run_best_effort([["brew", "install", "tesseract"]])

    if system == "linux" and command_exists("apt-get"):
        commands: List[List[str]] = []
        if hasattr(os, "geteuid") and os.geteuid() == 0:
            commands = [
                ["apt-get", "update"],
                ["apt-get", "install", "-y", "tesseract-ocr", "tesseract-ocr-chi-sim", "tesseract-ocr-chi-tra"],
            ]
        elif command_exists("sudo"):
            commands = [
                ["sudo", "-n", "apt-get", "update"],
                ["sudo", "-n", "apt-get", "install", "-y", "tesseract-ocr", "tesseract-ocr-chi-sim", "tesseract-ocr-chi-tra"],
            ]
        return run_best_effort(commands)

    if system == "windows":
        if command_exists("winget"):
            return run_best_effort([["winget", "install", "--id", "UB-Mannheim.TesseractOCR", "--silent", "--accept-package-agreements", "--accept-source-agreements"]])
        if command_exists("choco"):
            return run_best_effort([["choco", "install", "tesseract", "-y"]])

    print("Could not find a supported automatic installer for the Tesseract binary.", file=sys.stderr)
    return False


def run_best_effort(commands: Sequence[Sequence[str]]) -> bool:
    if not commands:
        return False
    for cmd in commands:
        try:
            result = run_command(list(cmd), check=False)
            if result.returncode != 0:
                return False
        except Exception as exc:
            print(f"Warning: command failed: {' '.join(cmd)}: {exc}", file=sys.stderr)
            return False
    return True


def ensure_local_ocr(auto_install: bool) -> bool:
    if local_ocr_available():
        return True

    if not python_package_installed("pytesseract"):
        print("pytesseract is missing.")
        if not auto_install:
            return False
        pip_install("pytesseract>=0.3.10")

    if not command_exists("tesseract"):
        print("Tesseract binary is missing.")
        if not auto_install:
            return False
        install_tesseract_binary()

    ok = local_ocr_available()
    if not ok:
        print("Local OCR is still unavailable after installation attempt.", file=sys.stderr)
    return ok


def load_dotenv_value(key: str, env_path: Path) -> Optional[str]:
    if not env_path.exists():
        return None
    for line in env_path.read_text(encoding="utf-8").splitlines():
        stripped = line.strip()
        if not stripped or stripped.startswith("#") or "=" not in stripped:
            continue
        name, value = stripped.split("=", 1)
        if name.strip() == key:
            return value.strip().strip('"').strip("'")
    return None


def save_dotenv_value(key: str, value: str, env_path: Path) -> None:
    lines: List[str] = []
    if env_path.exists():
        lines = env_path.read_text(encoding="utf-8").splitlines()
    replaced = False
    out_lines: List[str] = []
    for line in lines:
        if line.strip().startswith(f"{key}="):
            out_lines.append(f'{key}="{value}"')
            replaced = True
        else:
            out_lines.append(line)
    if not replaced:
        out_lines.append(f'{key}="{value}"')
    env_path.write_text("\n".join(out_lines) + "\n", encoding="utf-8")
    try:
        env_path.chmod(0o600)
    except Exception:
        pass


def dotenv_values(env_path: Path) -> Dict[str, str]:
    values: Dict[str, str] = {}
    if not env_path.exists():
        return values
    for line in env_path.read_text(encoding="utf-8").splitlines():
        stripped = line.strip()
        if not stripped or stripped.startswith("#") or "=" not in stripped:
            continue
        name, value = stripped.split("=", 1)
        values[name.strip()] = value.strip().strip('"').strip("'")
    return values


def save_dotenv_values(values: Dict[str, str], env_path: Path) -> None:
    existing_lines: List[str] = []
    if env_path.exists():
        existing_lines = env_path.read_text(encoding="utf-8").splitlines()
    replacements = {k: f'{k}="{v}"' for k, v in values.items() if v is not None}
    written: set[str] = set()
    out_lines: List[str] = []
    for line in existing_lines:
        stripped = line.strip()
        if "=" in stripped and not stripped.startswith("#"):
            name = stripped.split("=", 1)[0].strip()
            if name in replacements:
                out_lines.append(replacements[name])
                written.add(name)
                continue
        out_lines.append(line)
    for name, line in replacements.items():
        if name not in written:
            out_lines.append(line)
    env_path.write_text("\n".join(out_lines) + "\n", encoding="utf-8")
    try:
        env_path.chmod(0o600)
    except Exception:
        pass


def first_present(*values: Optional[str]) -> Optional[str]:
    for value in values:
        if value is not None and str(value).strip():
            return str(value).strip()
    return None


def choose_api_provider_interactively(default_provider: str) -> str:
    print("\nChoose an API text-recognition provider:")
    print("  1) OpenAI official API")
    print("  2) Third-party OpenAI-compatible API / gateway / proxy")
    print("  3) Cancel API OCR")
    while True:
        choice = input(f"Select 1, 2, or 3 [{ '1' if default_provider == 'openai' else '2' }]: ").strip()
        if not choice:
            return default_provider
        if choice == "1":
            return "openai"
        if choice == "2":
            return "openai-compatible"
        if choice == "3":
            return "cancel"
        print("Please enter 1, 2, or 3.")


def choose_endpoint_mode_interactively(default_mode: str) -> str:
    print("\nChoose the API call style:")
    print("  1) Auto: try Responses API first, then Chat Completions")
    print("  2) Responses API")
    print("  3) Chat Completions API (most third-party compatible services use this)")
    mapping = {"1": "auto", "2": "responses", "3": "chat-completions"}
    while True:
        choice = input("Select 1, 2, or 3 [1]: ").strip() or "1"
        if choice in mapping:
            return mapping[choice]
        print("Please enter 1, 2, or 3.")



def normalize_dedicated_ocr_provider(provider: str) -> str:
    provider = provider.lower().strip().replace("_", "-")
    aliases = {
        "ocrspace": "ocr-space",
        "ocr.space": "ocr-space",
        "azure": "azure-vision",
        "azure-read": "azure-vision",
        "azure-ai-vision": "azure-vision",
        "google": "google-vision",
        "google-cloud-vision": "google-vision",
        "gcv": "google-vision",
    }
    return aliases.get(provider, provider)


def dedicated_ocr_key_env(provider: str) -> str:
    if provider == "ocr-space":
        return "OCR_SPACE_API_KEY"
    if provider == "azure-vision":
        return "AZURE_VISION_KEY"
    if provider == "google-vision":
        return "GOOGLE_VISION_API_KEY"
    return "PPT_REBUILDER_OCR_API_KEY"


def choose_dedicated_ocr_provider_interactively(default_provider: str) -> str:
    print("\nChoose a dedicated OCR API provider:")
    print("  1) OCR.space API (simple OCR service; easiest to configure)")
    print("  2) Azure AI Vision Read OCR")
    print("  3) Google Cloud Vision OCR")
    print("  4) Cancel dedicated OCR API")
    default_choice = {"ocr-space": "1", "azure-vision": "2", "google-vision": "3"}.get(default_provider, "1")
    mapping = {"1": "ocr-space", "2": "azure-vision", "3": "google-vision", "4": "cancel"}
    while True:
        choice = input(f"Select 1, 2, 3, or 4 [{default_choice}]: ").strip() or default_choice
        if choice in mapping:
            return mapping[choice]
        print("Please enter 1, 2, 3, or 4.")


def ensure_dedicated_ocr_config(args: argparse.Namespace) -> Optional[DedicatedOcrConfig]:
    cached = getattr(args, "_dedicated_ocr_config_cache", None)
    if cached is not None:
        return cached

    env_path = Path(args.env_file).expanduser().resolve() if args.env_file else Path.cwd() / ".env"
    env = dotenv_values(env_path)

    provider = first_present(
        args.ocr_api_provider,
        os.environ.get("PPT_REBUILDER_OCR_API_PROVIDER"),
        env.get("PPT_REBUILDER_OCR_API_PROVIDER"),
        DEFAULT_OCR_API_PROVIDER,
    ) or DEFAULT_OCR_API_PROVIDER
    provider = normalize_dedicated_ocr_provider(provider)
    if provider not in {"ocr-space", "azure-vision", "google-vision"}:
        print(
            f"Unsupported dedicated OCR API provider '{provider}'. Use ocr-space, azure-vision, or google-vision.",
            file=sys.stderr,
        )
        return None

    if args.ocr_api_config_wizard and args.interactive:
        selected = choose_dedicated_ocr_provider_interactively(provider)
        if selected == "cancel":
            return None
        provider = selected

    default_key_env = dedicated_ocr_key_env(provider)
    key_env = args.ocr_api_key_env or default_key_env
    api_key = first_present(
        args.ocr_api_key,
        os.environ.get(key_env),
        env.get(key_env),
        os.environ.get("PPT_REBUILDER_OCR_API_KEY"),
        env.get("PPT_REBUILDER_OCR_API_KEY"),
        os.environ.get(default_key_env),
        env.get(default_key_env),
    )

    endpoint_default = DEFAULT_OCR_API_ENDPOINT if provider == "ocr-space" else None
    endpoint = first_present(
        args.ocr_api_endpoint,
        os.environ.get("PPT_REBUILDER_OCR_API_ENDPOINT"),
        env.get("PPT_REBUILDER_OCR_API_ENDPOINT"),
        os.environ.get("AZURE_VISION_ENDPOINT") if provider == "azure-vision" else None,
        env.get("AZURE_VISION_ENDPOINT") if provider == "azure-vision" else None,
        endpoint_default,
    )

    language = first_present(
        args.ocr_api_language,
        os.environ.get("PPT_REBUILDER_OCR_API_LANGUAGE"),
        env.get("PPT_REBUILDER_OCR_API_LANGUAGE"),
    )

    if args.interactive and args.ocr_api_config_wizard:
        if provider == "ocr-space":
            print("\nOCR.space selected. Ask your provider account page for the OCR.space API key.")
            if not endpoint:
                endpoint = DEFAULT_OCR_API_ENDPOINT
        elif provider == "azure-vision":
            print("\nAzure AI Vision selected. You need the Azure Vision endpoint URL and key.")
            print("Example endpoint format: https://YOUR-RESOURCE.cognitiveservices.azure.com")
            if not endpoint:
                endpoint = input("Azure Vision endpoint URL: ").strip()
        elif provider == "google-vision":
            print("\nGoogle Cloud Vision selected. This simple setup uses an API key for the REST API.")
            endpoint = endpoint or "https://vision.googleapis.com/v1/images:annotate"

        if not api_key:
            label = {
                "ocr-space": "OCR.space API key",
                "azure-vision": "Azure Vision key",
                "google-vision": "Google Vision API key",
            }.get(provider, "OCR API key")
            api_key = getpass.getpass(f"Enter {label}: ").strip()

        if provider == "ocr-space":
            current_lang = language or "eng"
            entered_lang = input(f"OCR.space language code [{current_lang}]: ").strip()
            if entered_lang:
                language = entered_lang
        elif provider == "azure-vision":
            current_lang = language or "auto"
            entered_lang = input(f"Azure OCR language hint, or auto [{current_lang}]: ").strip()
            if entered_lang and entered_lang.lower() != "auto":
                language = entered_lang
            elif not entered_lang:
                language = None
        elif provider == "google-vision":
            current_lang = language or "auto"
            entered_lang = input(f"Google OCR language hint, or auto [{current_lang}]: ").strip()
            if entered_lang and entered_lang.lower() != "auto":
                language = entered_lang
            elif not entered_lang:
                language = None

    if not api_key:
        print("No dedicated OCR API key provided; dedicated OCR API cannot run.", file=sys.stderr)
        return None
    if provider == "azure-vision" and not endpoint:
        print("Azure Vision OCR requires --ocr-api-endpoint or interactive endpoint entry.", file=sys.stderr)
        return None
    if provider == "ocr-space" and not endpoint:
        endpoint = DEFAULT_OCR_API_ENDPOINT
    if provider == "google-vision" and not endpoint:
        endpoint = "https://vision.googleapis.com/v1/images:annotate"

    os.environ[key_env] = api_key
    os.environ["PPT_REBUILDER_OCR_API_PROVIDER"] = provider
    if endpoint:
        os.environ["PPT_REBUILDER_OCR_API_ENDPOINT"] = endpoint
    if language:
        os.environ["PPT_REBUILDER_OCR_API_LANGUAGE"] = language

    save_values: Dict[str, str] = {}
    if args.interactive and (args.ocr_api_config_wizard or args.save_ocr_api_config):
        print(f"\nConfiguration file: {env_path}")
        if args.save_ocr_api_config or ask_yes_no("Save non-secret OCR API settings so next run can reuse them?", default=True):
            save_values["PPT_REBUILDER_OCR_API_PROVIDER"] = provider
            if endpoint:
                save_values["PPT_REBUILDER_OCR_API_ENDPOINT"] = endpoint
            if language:
                save_values["PPT_REBUILDER_OCR_API_LANGUAGE"] = language
        if args.save_ocr_api_key or ask_yes_no("Save the OCR API key to this local .env file? Only choose yes on your own private computer.", default=False):
            save_values[key_env] = api_key
            save_values["PPT_REBUILDER_OCR_API_KEY"] = api_key
        if save_values:
            save_dotenv_values(save_values, env_path)
            print(f"Saved dedicated OCR API settings to {env_path} with restricted file permissions when supported.")
    elif args.save_ocr_api_config or args.save_ocr_api_key:
        save_values["PPT_REBUILDER_OCR_API_PROVIDER"] = provider
        if endpoint:
            save_values["PPT_REBUILDER_OCR_API_ENDPOINT"] = endpoint
        if language:
            save_values["PPT_REBUILDER_OCR_API_LANGUAGE"] = language
        if args.save_ocr_api_key:
            save_values[key_env] = api_key
            save_values["PPT_REBUILDER_OCR_API_KEY"] = api_key
        save_dotenv_values(save_values, env_path)

    config = DedicatedOcrConfig(provider=provider, api_key=api_key, endpoint=endpoint, language=language)
    setattr(args, "_dedicated_ocr_config_cache", config)
    return config

def ensure_api_config(args: argparse.Namespace) -> Optional[ApiConfig]:
    cached = getattr(args, "_api_config_cache", None)
    if cached is not None:
        return cached
    env_path = Path(args.env_file).expanduser().resolve() if args.env_file else Path.cwd() / ".env"
    env = dotenv_values(env_path)

    provider = first_present(args.api_provider, os.environ.get("PPT_REBUILDER_API_PROVIDER"), env.get("PPT_REBUILDER_API_PROVIDER"), DEFAULT_API_PROVIDER) or DEFAULT_API_PROVIDER
    provider = provider.lower().strip()
    if provider in {"compatible", "third-party", "thirdparty", "custom"}:
        provider = "openai-compatible"
    if provider not in {"openai", "openai-compatible"}:
        print(f"Unsupported API provider '{provider}'. Use 'openai' or 'openai-compatible'.", file=sys.stderr)
        return None

    if args.api_config_wizard and args.interactive:
        selected = choose_api_provider_interactively(provider)
        if selected == "cancel":
            return None
        provider = selected

    key_env = args.api_key_env
    if not key_env:
        key_env = "OPENAI_API_KEY" if provider == "openai" else "PPT_REBUILDER_API_KEY"

    api_key = first_present(
        args.api_key,
        os.environ.get(key_env),
        env.get(key_env),
        os.environ.get("PPT_REBUILDER_API_KEY"),
        env.get("PPT_REBUILDER_API_KEY"),
        os.environ.get("OPENAI_API_KEY") if provider == "openai" else None,
        env.get("OPENAI_API_KEY") if provider == "openai" else None,
    )

    base_url = first_present(
        args.api_base_url,
        os.environ.get("PPT_REBUILDER_API_BASE_URL"),
        env.get("PPT_REBUILDER_API_BASE_URL"),
        os.environ.get("OPENAI_BASE_URL"),
        env.get("OPENAI_BASE_URL"),
    )
    if provider == "openai" and not args.api_base_url and not env.get("PPT_REBUILDER_API_BASE_URL") and not os.environ.get("PPT_REBUILDER_API_BASE_URL"):
        # Official OpenAI users normally do not need a custom base URL. Keep OPENAI_BASE_URL only if explicitly set.
        base_url = first_present(args.api_base_url, os.environ.get("OPENAI_BASE_URL"), env.get("OPENAI_BASE_URL"))

    model = first_present(
        args.api_model,
        os.environ.get("PPT_REBUILDER_API_MODEL"),
        env.get("PPT_REBUILDER_API_MODEL"),
        DEFAULT_API_MODEL,
    ) or DEFAULT_API_MODEL

    endpoint_mode = first_present(
        args.api_endpoint_mode,
        os.environ.get("PPT_REBUILDER_API_ENDPOINT_MODE"),
        env.get("PPT_REBUILDER_API_ENDPOINT_MODE"),
        DEFAULT_API_ENDPOINT_MODE,
    ) or DEFAULT_API_ENDPOINT_MODE
    endpoint_mode = endpoint_mode.lower().strip()
    if endpoint_mode in {"chat", "chat-completion", "chat_completion", "chat_completions"}:
        endpoint_mode = "chat-completions"
    if endpoint_mode not in {"auto", "responses", "chat-completions"}:
        print(f"Unsupported API endpoint mode '{endpoint_mode}'. Use auto, responses, or chat-completions.", file=sys.stderr)
        return None

    if args.interactive and provider == "openai-compatible" and not base_url:
        print("\nThird-party OpenAI-compatible API selected.")
        print("Ask your provider for its API Base URL, usually ending in /v1.")
        base_url = input("API Base URL: ").strip()

    if args.interactive and not api_key:
        label = "OPENAI_API_KEY" if provider == "openai" else "third-party API key"
        print(f"\nAPI OCR requires {label}.")
        api_key = getpass.getpass(f"Enter {label}: ").strip()

    if args.interactive and args.api_config_wizard:
        current_model = model or DEFAULT_API_MODEL
        entered_model = input(f"Model name [{current_model}]: ").strip()
        if entered_model:
            model = entered_model
        endpoint_mode = choose_endpoint_mode_interactively(endpoint_mode)

    if not api_key:
        print("No API key provided; API OCR cannot run.", file=sys.stderr)
        return None
    if provider == "openai-compatible" and not base_url:
        print("Third-party API mode requires --api-base-url or interactive Base URL entry.", file=sys.stderr)
        return None

    os.environ[key_env] = api_key
    if provider == "openai":
        os.environ["OPENAI_API_KEY"] = api_key
    if base_url:
        os.environ["PPT_REBUILDER_API_BASE_URL"] = base_url
    os.environ["PPT_REBUILDER_API_PROVIDER"] = provider
    os.environ["PPT_REBUILDER_API_MODEL"] = model
    os.environ["PPT_REBUILDER_API_ENDPOINT_MODE"] = endpoint_mode

    save_values: Dict[str, str] = {}
    if args.interactive and (args.api_config_wizard or args.save_api_config):
        print(f"\nConfiguration file: {env_path}")
        if args.save_api_config or ask_yes_no("Save non-secret API settings so next run can reuse them?", default=True):
            save_values.update(
                {
                    "PPT_REBUILDER_API_PROVIDER": provider,
                    "PPT_REBUILDER_API_MODEL": model,
                    "PPT_REBUILDER_API_ENDPOINT_MODE": endpoint_mode,
                }
            )
            if base_url:
                save_values["PPT_REBUILDER_API_BASE_URL"] = base_url
        if args.save_api_key or ask_yes_no("Save the API key to this local .env file? Only choose yes on your own private computer.", default=False):
            save_values[key_env] = api_key
            if provider == "openai":
                save_values["OPENAI_API_KEY"] = api_key
            else:
                save_values["PPT_REBUILDER_API_KEY"] = api_key
        if save_values:
            save_dotenv_values(save_values, env_path)
            print(f"Saved API settings to {env_path} with restricted file permissions when supported.")
    elif args.save_api_config or args.save_api_key:
        save_values.update(
            {
                "PPT_REBUILDER_API_PROVIDER": provider,
                "PPT_REBUILDER_API_MODEL": model,
                "PPT_REBUILDER_API_ENDPOINT_MODE": endpoint_mode,
            }
        )
        if base_url:
            save_values["PPT_REBUILDER_API_BASE_URL"] = base_url
        if args.save_api_key:
            save_values[key_env] = api_key
            if provider == "openai":
                save_values["OPENAI_API_KEY"] = api_key
            else:
                save_values["PPT_REBUILDER_API_KEY"] = api_key
        save_dotenv_values(save_values, env_path)

    if not python_package_installed("openai"):
        print("openai Python package is missing.")
        if args.auto_install_api_package:
            pip_install("openai>=1.0.0")
        else:
            print("Install it with: pip install openai", file=sys.stderr)
    if not python_package_installed("openai"):
        return None

    config = ApiConfig(provider=provider, api_key=api_key, model=model, base_url=base_url, endpoint_mode=endpoint_mode)
    setattr(args, "_api_config_cache", config)
    return config

def ask_yes_no(question: str, default: bool = False) -> bool:
    suffix = "[Y/n]" if default else "[y/N]"
    answer = input(f"{question} {suffix} ").strip().lower()
    if not answer:
        return default
    return answer in {"y", "yes", "1", "true", "是", "好"}


def choose_ocr_mode(args: argparse.Namespace) -> str:
    if args.ocr == "none":
        return "none"
    if args.ocr == "json":
        if not args.ocr_json:
            raise SystemExit("--ocr json requires --ocr-json /path/to/ocr.json")
        return "json"
    if args.ocr == "local":
        if ensure_local_ocr(auto_install=args.auto_install_local):
            return "local"
        raise SystemExit("Local OCR is unavailable. Rerun with --ocr api or --ocr none.")
    if args.ocr == "api":
        if args.interactive and not args.api_provider:
            args.api_config_wizard = True
        if ensure_api_config(args):
            return "api"
        raise SystemExit("AI vision recognition is unavailable. Provide API settings or rerun with --ocr local/ocr-api/none.")
    if args.ocr == "ocr-api":
        if args.interactive and not args.ocr_api_provider:
            args.ocr_api_config_wizard = True
        if ensure_dedicated_ocr_config(args):
            return "ocr-api"
        raise SystemExit("Dedicated OCR API is unavailable. Provide OCR API settings or rerun with --ocr local/api/none.")

    # auto mode: prefer installed local OCR; if missing, ask the user what to do.
    if local_ocr_available():
        return "local"
    if args.ocr_json:
        return "json"

    if not args.interactive:
        print("Local OCR is not available and this is non-interactive; continuing without OCR.", file=sys.stderr)
        return "none"

    while True:
        print("\nLocal OCR is not available. Choose how to continue:")
        print("  1) Install local OCR automatically, then use local OCR")
        print("  2) Configure AI vision recognition: OpenAI or third-party OpenAI-compatible API")
        print("  3) Configure dedicated OCR API: OCR.space, Azure Vision, or Google Vision")
        print("  4) Skip recognition and create a visual-reference PPT only")
        choice = input("Select 1, 2, 3, or 4: ").strip()
        if choice == "1":
            if ensure_local_ocr(auto_install=True):
                return "local"
            print("Automatic local OCR installation did not finish successfully.")
        elif choice == "2":
            args.api_config_wizard = True
            if ensure_api_config(args):
                return "api"
            print("AI vision configuration did not finish successfully.")
        elif choice == "3":
            args.ocr_api_config_wizard = True
            if ensure_dedicated_ocr_config(args):
                return "ocr-api"
            print("Dedicated OCR API configuration did not finish successfully.")
        elif choice == "4":
            return "none"
        else:
            print("Please enter 1, 2, 3, or 4.")


def ocr_with_tesseract(image_path: Path, min_confidence: float, lang: str) -> List[TextItem]:
    try:
        import pytesseract
        from pytesseract import Output
    except ImportError:
        print("Warning: pytesseract is not installed; skipping local OCR.", file=sys.stderr)
        return []

    try:
        with Image.open(image_path) as img:
            width_px, height_px = img.size
            kwargs = {"output_type": Output.DICT}
            if lang:
                kwargs["lang"] = lang
            data = pytesseract.image_to_data(img, **kwargs)
    except Exception as exc:
        print(f"Warning: local OCR failed for {image_path.name}: {exc}", file=sys.stderr)
        return []

    grouped: Dict[Tuple[int, int, int], List[int]] = {}
    n = len(data.get("text", []))
    for i in range(n):
        text = str(data["text"][i]).strip()
        conf = parse_conf(data.get("conf", [])[i])
        if not text or conf < min_confidence:
            continue
        key = (
            int(data.get("block_num", [0] * n)[i]),
            int(data.get("par_num", [0] * n)[i]),
            int(data.get("line_num", [0] * n)[i]),
        )
        grouped.setdefault(key, []).append(i)

    items: List[TextItem] = []
    for _, indexes in sorted(grouped.items()):
        texts = [str(data["text"][i]).strip() for i in indexes]
        line_text = " ".join(t for t in texts if t)
        if not line_text:
            continue
        lefts = [int(data["left"][i]) for i in indexes]
        tops = [int(data["top"][i]) for i in indexes]
        rights = [int(data["left"][i]) + int(data["width"][i]) for i in indexes]
        bottoms = [int(data["top"][i]) + int(data["height"][i]) for i in indexes]
        x1, y1, x2, y2 = min(lefts), min(tops), max(rights), max(bottoms)
        confs = [parse_conf(data["conf"][i]) for i in indexes]
        items.append(
            TextItem(
                text=line_text,
                x=x1 / width_px,
                y=y1 / height_px,
                w=max(1, x2 - x1) / width_px,
                h=max(1, y2 - y1) / height_px,
                confidence=sum(confs) / max(len(confs), 1),
            )
        )
    return items


def data_url_for_image(image_path: Path) -> str:
    mime = mimetypes.guess_type(str(image_path))[0] or "image/png"
    with image_path.open("rb") as f:
        encoded = base64.b64encode(f.read()).decode("utf-8")
    return f"data:{mime};base64,{encoded}"


def extract_json_object(text: str) -> Dict[str, Any]:
    text = text.strip()
    if text.startswith("```"):
        text = re.sub(r"^```(?:json)?", "", text).strip()
        text = re.sub(r"```$", "", text).strip()
    try:
        return json.loads(text)
    except json.JSONDecodeError:
        start = text.find("{")
        end = text.rfind("}")
        if start >= 0 and end > start:
            return json.loads(text[start : end + 1])
        raise


def api_ocr_prompt() -> str:
    return (
        "You are converting a PowerPoint slide screenshot back into editable PPT text. "
        "Extract visible slide text and estimate each text block's bounding box. "
        "Return ONLY valid JSON with this shape: "
        "{\"items\":[{\"text\":\"...\",\"x\":0.0,\"y\":0.0,\"w\":0.1,\"h\":0.05,\"font_size\":18,\"confidence\":90}]}. "
        "Use normalized coordinates from 0 to 1 relative to the image. "
        "Group nearby words into natural editable lines or text boxes. "
        "Preserve the original language and punctuation. Do not describe non-text graphics."
    )


def make_openai_client(api_config: ApiConfig) -> Any:
    from openai import OpenAI

    kwargs: Dict[str, Any] = {"api_key": api_config.api_key}
    if api_config.base_url:
        kwargs["base_url"] = api_config.base_url
    return OpenAI(**kwargs)


def parse_api_text_to_items(output_text: str) -> List[TextItem]:
    payload = extract_json_object(output_text)
    return text_items_from_raw(payload.get("items", []))


def ocr_with_responses_api(client: Any, image_path: Path, model: str, prompt: str) -> List[TextItem]:
    response = client.responses.create(
        model=model,
        input=[
            {
                "role": "user",
                "content": [
                    {"type": "input_text", "text": prompt},
                    {"type": "input_image", "image_url": data_url_for_image(image_path)},
                ],
            }
        ],
    )
    return parse_api_text_to_items(response.output_text)


def ocr_with_chat_completions_api(client: Any, image_path: Path, model: str, prompt: str) -> List[TextItem]:
    messages = [
        {
            "role": "user",
            "content": [
                {"type": "text", "text": prompt},
                {"type": "image_url", "image_url": {"url": data_url_for_image(image_path)}},
            ],
        }
    ]
    try:
        response = client.chat.completions.create(
            model=model,
            messages=messages,
            response_format={"type": "json_object"},
        )
    except Exception:
        # Many OpenAI-compatible third-party gateways do not support response_format.
        response = client.chat.completions.create(model=model, messages=messages)
    text = response.choices[0].message.content or ""
    return parse_api_text_to_items(text)


def ocr_with_api(image_path: Path, api_config: ApiConfig, prompt: Optional[str] = None) -> List[TextItem]:
    try:
        import openai  # noqa: F401
    except ImportError:
        print("Warning: openai package is not installed; skipping API OCR.", file=sys.stderr)
        return []

    final_prompt = prompt or api_ocr_prompt()
    try:
        client = make_openai_client(api_config)
        modes = [api_config.endpoint_mode]
        if api_config.endpoint_mode == "auto":
            modes = ["responses", "chat-completions"]
        last_error: Optional[Exception] = None
        for mode in modes:
            try:
                if mode == "responses":
                    return ocr_with_responses_api(client, image_path, api_config.model, final_prompt)
                if mode == "chat-completions":
                    return ocr_with_chat_completions_api(client, image_path, api_config.model, final_prompt)
            except Exception as exc:
                last_error = exc
                if api_config.endpoint_mode != "auto":
                    raise
        if last_error:
            raise last_error
    except Exception as exc:
        location = api_config.base_url or "OpenAI default endpoint"
        print(f"Warning: API OCR failed for {image_path.name} using {api_config.provider} at {location}: {exc}", file=sys.stderr)
    return []


def http_request_json(url: str, *, headers: Dict[str, str], body: bytes) -> Dict[str, Any]:
    request = urllib.request.Request(url, data=body, headers=headers, method="POST")
    try:
        with urllib.request.urlopen(request, timeout=90) as response:
            raw = response.read().decode("utf-8", errors="replace")
    except urllib.error.HTTPError as exc:
        raw_error = exc.read().decode("utf-8", errors="replace")
        raise RuntimeError(f"HTTP {exc.code}: {raw_error[:1000]}") from exc
    return json.loads(raw)


def box_from_points(points: Sequence[Any], image_width: int, image_height: int) -> Optional[Tuple[float, float, float, float]]:
    xs: List[float] = []
    ys: List[float] = []
    if points and all(isinstance(point, (int, float)) for point in points):
        flat = list(points)
        for i in range(0, len(flat) - 1, 2):
            xs.append(float(flat[i]))
            ys.append(float(flat[i + 1]))
    else:
        for point in points:
            if isinstance(point, dict):
                if "x" in point and "y" in point:
                    xs.append(float(point.get("x", 0)))
                    ys.append(float(point.get("y", 0)))
            elif isinstance(point, (list, tuple)) and len(point) >= 2:
                xs.append(float(point[0]))
                ys.append(float(point[1]))
    if not xs or not ys:
        return None
    x1, x2 = min(xs), max(xs)
    y1, y2 = min(ys), max(ys)
    return x1 / max(image_width, 1), y1 / max(image_height, 1), max(1.0, x2 - x1) / max(image_width, 1), max(1.0, y2 - y1) / max(image_height, 1)


def add_text_item_from_box(items: List[TextItem], text: str, box: Optional[Tuple[float, float, float, float]], confidence: Optional[float] = None) -> None:
    clean = str(text or "").strip()
    if not clean or not box:
        return
    x, y, w, h = box
    items.append(TextItem(text=clean, x=x, y=y, w=w, h=h, confidence=confidence))


def ocr_with_ocr_space(image_path: Path, config: DedicatedOcrConfig) -> List[TextItem]:
    endpoint = config.endpoint or DEFAULT_OCR_API_ENDPOINT
    language = config.language or "eng"
    form = {
        "apikey": config.api_key,
        "language": language,
        "isOverlayRequired": "true",
        "scale": "true",
        "OCREngine": "2",
        "base64Image": data_url_for_image(image_path),
    }
    body = urllib.parse.urlencode(form).encode("utf-8")
    payload = http_request_json(
        endpoint,
        headers={"Content-Type": "application/x-www-form-urlencoded"},
        body=body,
    )
    if payload.get("IsErroredOnProcessing"):
        error = payload.get("ErrorMessage") or payload.get("ErrorDetails") or "unknown OCR.space error"
        raise RuntimeError(str(error))

    with Image.open(image_path) as img:
        image_width, image_height = img.size

    items: List[TextItem] = []
    for result in payload.get("ParsedResults", []) or []:
        overlay = result.get("TextOverlay") or {}
        for line in overlay.get("Lines", []) or []:
            words = line.get("Words", []) or []
            text = line.get("LineText") or " ".join(str(word.get("WordText", "")).strip() for word in words)
            if words:
                lefts = [float(word.get("Left", 0)) for word in words]
                tops = [float(word.get("Top", 0)) for word in words]
                rights = [float(word.get("Left", 0)) + float(word.get("Width", 0)) for word in words]
                bottoms = [float(word.get("Top", 0)) + float(word.get("Height", 0)) for word in words]
                box = (min(lefts) / image_width, min(tops) / image_height, max(1.0, max(rights) - min(lefts)) / image_width, max(1.0, max(bottoms) - min(tops)) / image_height)
                add_text_item_from_box(items, text, box)
        if not items:
            parsed = str(result.get("ParsedText", "")).strip()
            if parsed:
                for i, line_text in enumerate([line.strip() for line in parsed.splitlines() if line.strip()]):
                    add_text_item_from_box(items, line_text, (0.05, 0.05 + i * 0.06, 0.90, 0.05))
    return items


def azure_vision_url(endpoint: str, language: Optional[str]) -> str:
    endpoint = endpoint.rstrip("/")
    if "imageanalysis:analyze" in endpoint:
        url = endpoint
    else:
        url = f"{endpoint}/computervision/imageanalysis:analyze"
    sep = "&" if "?" in url else "?"
    params = {"api-version": DEFAULT_AZURE_VISION_API_VERSION, "features": "read"}
    if language:
        params["language"] = language
    return url + sep + urllib.parse.urlencode(params)


def ocr_with_azure_vision(image_path: Path, config: DedicatedOcrConfig) -> List[TextItem]:
    if not config.endpoint:
        raise RuntimeError("Azure Vision endpoint is required.")
    with Image.open(image_path) as img:
        image_width, image_height = img.size
    mime = mimetypes.guess_type(str(image_path))[0] or "image/png"
    body = image_path.read_bytes()
    payload = http_request_json(
        azure_vision_url(config.endpoint, config.language),
        headers={"Ocp-Apim-Subscription-Key": config.api_key, "Content-Type": mime},
        body=body,
    )
    items: List[TextItem] = []
    read_result = payload.get("readResult") or payload.get("read_result") or {}
    for block in read_result.get("blocks", []) or []:
        for line in block.get("lines", []) or []:
            box = box_from_points(line.get("boundingPolygon") or line.get("bounding_poly") or [], image_width, image_height)
            add_text_item_from_box(items, line.get("text", ""), box)
    return items


def google_vision_url(endpoint: str, api_key: str) -> str:
    endpoint = endpoint or "https://vision.googleapis.com/v1/images:annotate"
    if "key=" in endpoint:
        return endpoint
    sep = "&" if "?" in endpoint else "?"
    return endpoint + sep + urllib.parse.urlencode({"key": api_key})


def ocr_with_google_vision(image_path: Path, config: DedicatedOcrConfig) -> List[TextItem]:
    with Image.open(image_path) as img:
        image_width, image_height = img.size
    encoded = base64.b64encode(image_path.read_bytes()).decode("utf-8")
    image_context: Dict[str, Any] = {}
    if config.language:
        image_context["languageHints"] = [config.language]
    request_item: Dict[str, Any] = {
        "image": {"content": encoded},
        "features": [{"type": "DOCUMENT_TEXT_DETECTION"}],
    }
    if image_context:
        request_item["imageContext"] = image_context
    body = json.dumps({"requests": [request_item]}).encode("utf-8")
    payload = http_request_json(
        google_vision_url(config.endpoint or "https://vision.googleapis.com/v1/images:annotate", config.api_key),
        headers={"Content-Type": "application/json"},
        body=body,
    )
    response = (payload.get("responses") or [{}])[0]
    if response.get("error"):
        raise RuntimeError(json.dumps(response["error"], ensure_ascii=False))
    annotations = response.get("textAnnotations") or []
    items: List[TextItem] = []
    for annotation in annotations[1:]:
        vertices = ((annotation.get("boundingPoly") or {}).get("vertices") or [])
        box = box_from_points(vertices, image_width, image_height)
        add_text_item_from_box(items, annotation.get("description", ""), box)
    if not items and annotations:
        vertices = ((annotations[0].get("boundingPoly") or {}).get("vertices") or [])
        box = box_from_points(vertices, image_width, image_height) or (0.05, 0.05, 0.90, 0.90)
        add_text_item_from_box(items, annotations[0].get("description", ""), box)
    return items


def ocr_with_dedicated_api(image_path: Path, config: DedicatedOcrConfig) -> List[TextItem]:
    try:
        if config.provider == "ocr-space":
            return ocr_with_ocr_space(image_path, config)
        if config.provider == "azure-vision":
            return ocr_with_azure_vision(image_path, config)
        if config.provider == "google-vision":
            return ocr_with_google_vision(image_path, config)
        raise RuntimeError(f"Unsupported dedicated OCR API provider: {config.provider}")
    except Exception as exc:
        print(f"Warning: dedicated OCR API failed for {image_path.name} using {config.provider}: {exc}", file=sys.stderr)
        return []

def normalize_box(item: TextItem, image_width: int, image_height: int) -> Tuple[float, float, float, float]:
    values = [item.x, item.y, item.w, item.h]
    if all(-0.01 <= v <= 1.2 for v in values):
        return item.x, item.y, item.w, item.h
    return (
        item.x / max(image_width, 1),
        item.y / max(image_height, 1),
        item.w / max(image_width, 1),
        item.h / max(image_height, 1),
    )


def hex_to_rgb(value: str) -> RGBColor:
    value = value.strip().lstrip("#")
    if len(value) != 6 or not re.fullmatch(r"[0-9a-fA-F]{6}", value):
        raise argparse.ArgumentTypeError("Color must be a 6-digit hex value, such as 000000")
    return RGBColor(int(value[0:2], 16), int(value[2:4], 16), int(value[4:6], 16))


def add_text_item(slide: Any, item: TextItem, image_size: Tuple[int, int], slide_size: Tuple[int, int], text_color: RGBColor) -> None:
    image_width, image_height = image_size
    slide_width, slide_height = slide_size
    x, y, w, h = normalize_box(item, image_width, image_height)
    left = int(max(0, x) * slide_width)
    top = int(max(0, y) * slide_height)
    width = int(max(0.01, w) * slide_width)
    height = int(max(0.01, h) * slide_height)

    box = slide.shapes.add_textbox(left, top, width, height)
    box.name = "editable_ocr_text"
    text_frame = box.text_frame
    text_frame.clear()
    text_frame.margin_left = 0
    text_frame.margin_right = 0
    text_frame.margin_top = 0
    text_frame.margin_bottom = 0
    text_frame.word_wrap = False

    paragraph = text_frame.paragraphs[0]
    run = paragraph.add_run()
    run.text = item.text
    font = run.font
    font.name = "Arial"
    if item.font_size:
        size = max(5.0, min(96.0, item.font_size))
    else:
        box_height_inches = height / EMU_PER_INCH
        size = max(6.0, min(44.0, box_height_inches * 72 * 0.85))
    font.size = Pt(size)
    font.color.rgb = text_color

    try:
        box.fill.background()
        box.line.fill.background()
    except Exception:
        pass


def items_for_image(
    image_path: Path,
    index: int,
    ocr_mode: str,
    ocr_data: Dict[str, List[TextItem]],
    min_confidence: float,
    ocr_lang: str,
    api_model: str,
    api_config: Optional[ApiConfig],
    api_prompt: Optional[str],
    dedicated_ocr_config: Optional[DedicatedOcrConfig],
) -> List[TextItem]:
    if ocr_mode == "none":
        return []
    if ocr_mode == "json":
        return ocr_data.get(image_path.name) or ocr_data.get(str(index + 1)) or []
    if ocr_mode == "local":
        explicit = ocr_data.get(image_path.name) or ocr_data.get(str(index + 1)) or []
        if explicit:
            return explicit
        return ocr_with_tesseract(image_path, min_confidence, ocr_lang)
    if ocr_mode == "api":
        explicit = ocr_data.get(image_path.name) or ocr_data.get(str(index + 1)) or []
        if explicit:
            return explicit
        if not api_config:
            return []
        return ocr_with_api(image_path, api_config, api_prompt)
    if ocr_mode == "ocr-api":
        explicit = ocr_data.get(image_path.name) or ocr_data.get(str(index + 1)) or []
        if explicit:
            return explicit
        if not dedicated_ocr_config:
            return []
        return ocr_with_dedicated_api(image_path, dedicated_ocr_config)
    raise ValueError(f"Unsupported OCR mode: {ocr_mode}")


def build_pptx(args: argparse.Namespace) -> None:
    input_path = Path(args.input).expanduser().resolve()
    output_path = Path(args.output).expanduser().resolve()
    images = collect_images(input_path)

    slide_width, slide_height = infer_slide_size(images[0], args.slide_size)
    prs = Presentation()
    prs.slide_width = slide_width
    prs.slide_height = slide_height
    blank_layout = prs.slide_layouts[6]

    ocr_data: Dict[str, List[TextItem]] = {}
    if args.ocr_json:
        ocr_data = load_ocr_json(Path(args.ocr_json).expanduser().resolve())

    resolved_ocr_mode = choose_ocr_mode(args)
    api_config = ensure_api_config(args) if resolved_ocr_mode == "api" else None
    dedicated_ocr_config = ensure_dedicated_ocr_config(args) if resolved_ocr_mode == "ocr-api" else None
    api_prompt = Path(args.api_prompt).read_text(encoding="utf-8") if args.api_prompt else None

    total_text_items = 0
    with tempfile.TemporaryDirectory() as tmp:
        tmp_dir = Path(tmp)
        for index, image_path in enumerate(images):
            slide = prs.slides.add_slide(blank_layout)
            with Image.open(image_path) as img:
                image_width, image_height = img.size

            if args.background_mode != "none":
                pic_path = faded_copy(image_path, args.reference_opacity, tmp_dir)
                left, top, width, height = fit_rect(image_width, image_height, slide_width, slide_height)
                slide.shapes.add_picture(str(pic_path), left, top, width=width, height=height)

            text_items = items_for_image(
                image_path=image_path,
                index=index,
                ocr_mode=resolved_ocr_mode,
                ocr_data=ocr_data,
                min_confidence=args.min_confidence,
                ocr_lang=args.ocr_lang,
                api_model=args.api_model,
                api_config=api_config,
                api_prompt=api_prompt,
                dedicated_ocr_config=dedicated_ocr_config,
            )
            if args.max_text_items and len(text_items) > args.max_text_items:
                text_items = text_items[: args.max_text_items]
            total_text_items += len(text_items)
            for item in text_items:
                add_text_item(
                    slide=slide,
                    item=item,
                    image_size=(image_width, image_height),
                    slide_size=(slide_width, slide_height),
                    text_color=args.text_color,
                )

    output_path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(output_path)
    print(f"Created {output_path}")
    print(f"Slides: {len(images)}")
    print(f"Background mode: {args.background_mode}")
    print(f"OCR mode: {resolved_ocr_mode}")
    print(f"Editable text boxes: {total_text_items}")


def parse_args(argv: Optional[Iterable[str]] = None) -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Rebuild a best-effort editable PPTX from slide images.")
    parser.add_argument("--input", required=True, help="Input image file or directory of slide images.")
    parser.add_argument("--output", required=True, help="Output .pptx path.")
    parser.add_argument("--slide-size", choices=["auto", "16:9", "4:3"], default="auto", help="Output slide size.")
    parser.add_argument("--background-mode", choices=["full", "none"], default="full", help="Place source images on slides or omit them.")
    parser.add_argument("--reference-opacity", type=float, default=1.0, help="Opacity for source image layer, from 0 to 1.")
    parser.add_argument("--ocr", choices=["auto", "local", "api", "ocr-api", "json", "none"], default="auto", help="Recognition source for editable text boxes: local OCR, AI vision API, dedicated OCR API, JSON, or none.")
    parser.add_argument("--ocr-json", help="Optional sidecar OCR JSON file.")
    parser.add_argument("--ocr-lang", default="eng+chi_sim", help="Tesseract language setting, e.g. eng, chi_sim, or eng+chi_sim.")
    parser.add_argument("--min-confidence", type=float, default=55.0, help="Minimum local OCR confidence for Tesseract words.")
    parser.add_argument("--max-text-items", type=int, default=0, help="Optional cap on text boxes per slide; 0 means no cap.")
    parser.add_argument("--text-color", type=hex_to_rgb, default=hex_to_rgb("000000"), help="Editable text color as hex, default 000000.")
    parser.add_argument("--api-provider", choices=["openai", "openai-compatible"], help="API provider. Use openai-compatible for third-party providers, gateways, or proxies.")
    parser.add_argument("--api-base-url", help="Base URL for OpenAI-compatible third-party APIs, usually ending in /v1.")
    parser.add_argument("--api-endpoint-mode", choices=["auto", "responses", "chat-completions"], help="API call style. auto tries Responses API then Chat Completions.")
    parser.add_argument("--api-model", help=f"Vision model for --ocr api, default {DEFAULT_API_MODEL}. For third-party APIs, use the model name from that provider.")
    parser.add_argument("--api-key", help="API key. Prefer environment variables, .env, or interactive entry instead of this flag.")
    parser.add_argument("--api-key-env", help="Name of the environment variable that stores the API key. Defaults to OPENAI_API_KEY for OpenAI and PPT_REBUILDER_API_KEY for third-party APIs.")
    parser.add_argument("--env-file", help="Path to .env file for reading/saving API configuration. Default: ./.env")
    parser.add_argument("--api-prompt", help="Optional path to a custom prompt for API OCR JSON extraction.")
    parser.add_argument("--api-config-wizard", action="store_true", help="Ask beginner-friendly questions to configure OpenAI or third-party AI vision recognition.")
    parser.add_argument("--ocr-api-provider", choices=["ocr-space", "azure-vision", "google-vision"], help="Dedicated OCR API provider for --ocr ocr-api.")
    parser.add_argument("--ocr-api-endpoint", help="Dedicated OCR API endpoint. Required for Azure Vision; optional for OCR.space and Google Vision.")
    parser.add_argument("--ocr-api-key", help="Dedicated OCR API key. Prefer environment variables, .env, or interactive entry instead of this flag.")
    parser.add_argument("--ocr-api-key-env", help="Name of the environment variable that stores the dedicated OCR API key.")
    parser.add_argument("--ocr-api-language", help="Optional language hint/code for the dedicated OCR API, such as eng, chs, zh-Hans, or en.")
    parser.add_argument("--ocr-api-config-wizard", action="store_true", help="Ask beginner-friendly questions to configure a dedicated OCR API such as OCR.space, Azure Vision, or Google Vision.")
    parser.add_argument("--save-ocr-api-config", action="store_true", help="Save dedicated OCR provider/endpoint/language settings to the .env file without asking again.")
    parser.add_argument("--save-ocr-api-key", action="store_true", help="Save an interactively supplied dedicated OCR API key to the .env file without asking again.")
    parser.add_argument("--save-api-config", action="store_true", help="Save provider/base URL/model settings to the .env file without asking again.")
    parser.add_argument("--save-api-key", action="store_true", help="Save an interactively supplied API key to the .env file without asking again.")
    parser.add_argument("--auto-install-local", action="store_true", help="For --ocr local, attempt automatic local OCR installation if missing.")
    parser.add_argument("--auto-install-api-package", action="store_true", default=True, help="Install the openai Python package automatically when --ocr api needs it.")
    parser.add_argument("--non-interactive", dest="interactive", action="store_false", help="Do not prompt; degrade to no OCR in --ocr auto when local OCR is missing.")
    parser.set_defaults(interactive=True)
    return parser.parse_args(argv)


def main(argv: Optional[Iterable[str]] = None) -> int:
    args = parse_args(argv)
    build_pptx(args)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
